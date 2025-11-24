from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, date
from io import BytesIO
from typing import Iterable, List

import xlrd
from openpyxl import load_workbook


@dataclass
class SheetData:
    name: str
    rows: List[List[str]]
    has_merges: bool


class ExcelToMarkdownHandler:
    """负责解析 Excel 并生成 Markdown 输出（多 Sheet 支持）。"""

    def convert(self, file_bytes: bytes, ext: str) -> str:
        if not isinstance(file_bytes, (bytes, bytearray)):
            raise TypeError("file_bytes must be bytes")
        if not file_bytes:
            raise ValueError("Excel file is empty")

        ext = ext.lower()
        if ext == "xlsx":
            sheets = self._parse_xlsx(file_bytes)
        elif ext == "xls":
            sheets = self._parse_xls(file_bytes)
        else:
            raise ValueError(f"Unsupported excel extension: {ext}")

        rendered_parts: List[str] = []
        for sheet in sheets:
            if not sheet.rows:
                continue

            rendered_parts.append(f"## Sheet: {sheet.name}")

            if sheet.has_merges:
                rendered_parts.append(self._rows_to_code_block(sheet.rows))
            else:
                rendered_parts.append(self._rows_to_markdown(sheet.rows))

        if not rendered_parts:
            rendered_parts.append("（Excel 文件无有效数据）")

        return "\n\n".join(rendered_parts).strip()

    def _parse_xlsx(self, file_bytes: bytes) -> List[SheetData]:
        buffer = BytesIO(file_bytes)
        workbook = load_workbook(
            buffer,
            data_only=True,
        )

        sheets: List[SheetData] = []
        for worksheet in workbook.worksheets:
            rows = self._collect_rows(worksheet.iter_rows(values_only=True))
            merged_ranges = getattr(worksheet, "merged_cells", None)
            has_merges = bool(getattr(merged_ranges, "ranges", merged_ranges))
            sheets.append(SheetData(name=worksheet.title, rows=rows, has_merges=has_merges))

        workbook.close()

        return sheets

    def _parse_xls(self, file_bytes: bytes) -> List[SheetData]:
        workbook = xlrd.open_workbook(file_contents=file_bytes)

        sheets: List[SheetData] = []
        for sheet in workbook.sheets():
            rows: List[List[str]] = []
            for row_idx in range(sheet.nrows):
                row_values = [
                    self._format_value(sheet.cell_value(row_idx, col_idx))
                    for col_idx in range(sheet.ncols)
                ]
                if self._is_row_empty(row_values):
                    continue
                rows.append(row_values)

            has_merges = bool(getattr(sheet, "merged_cells", []))
            sheets.append(SheetData(name=sheet.name, rows=rows, has_merges=has_merges))

        return sheets

    def _collect_rows(self, rows_iter: Iterable[Iterable[object | None]]) -> List[List[str]]:
        rows: List[List[str]] = []
        for row in rows_iter:
            row_values = [self._format_value(cell) for cell in row]
            if self._is_row_empty(row_values):
                continue
            rows.append(row_values)
        return rows

    @staticmethod
    def _is_row_empty(row_values: List[str]) -> bool:
        return all(value == "" for value in row_values)

    def _rows_to_markdown(self, rows: List[List[str]]) -> str:
        if not rows:
            return ""

        max_columns = max(len(row) for row in rows)
        padded_rows = [self._pad_row(row, max_columns) for row in rows]

        header = padded_rows[0]
        body = padded_rows[1:]

        lines: List[str] = []
        lines.append(self._format_row(header))
        lines.append(self._format_row(["---"] * max_columns))

        for row in body:
            lines.append(self._format_row(row))

        if len(padded_rows) == 1:
            # 只有表头时仍需保留分隔线
            pass

        return "\n".join(lines)

    @staticmethod
    def _rows_to_code_block(rows: List[List[str]]) -> str:
        rendered_rows = ["\t".join(row) for row in rows]
        return "```\n" + "\n".join(rendered_rows) + "\n```"

    @staticmethod
    def _pad_row(row: List[str], length: int) -> List[str]:
        if len(row) >= length:
            return row
        return row + [""] * (length - len(row))

    @staticmethod
    def _format_row(row: List[str]) -> str:
        return "| " + " | ".join(row) + " |"

    @staticmethod
    def _format_value(value: object | None) -> str:
        if value is None:
            return ""
        if isinstance(value, str):
            return value.strip()
        if isinstance(value, (int, float)):
            if isinstance(value, float) and value.is_integer():
                return str(int(value))
            return str(value)
        if isinstance(value, (datetime, date)):
            return value.isoformat()
        return str(value)

